"""Two-sheet tracker: "Leads" (sales-facing) and "Technical" (run id, evidence,
provenance status, sanctions API, telemetry)."""
from pathlib import Path

from openpyxl import load_workbook

from leadscout.models import CloudUsageAssessment, ComplianceResult, FitResult, Lead, LeadOutcome, Research
from leadscout.tracker import LEADS_COLUMNS, TECHNICAL_COLUMNS, write_row


def test_leads_sheet_has_sales_facing_columns():
    names = [name for name, _ in LEADS_COLUMNS]
    assert "Fit score" in names
    assert "Fit confidence" in names
    assert "Spend prior (relative)" in names
    assert "Compliance" in names
    assert "Cloud usage" in names
    assert "Sales-ready" in names
    # Technical/internal details do not belong on the sales-facing sheet.
    assert "Run ID" not in names
    assert "LLM calls" not in names


def test_technical_sheet_has_run_and_telemetry_columns():
    names = [name for name, _ in TECHNICAL_COLUMNS]
    for expected in ("Run ID", "Evidence count", "Provenance status", "Provenance head",
                     "Sanctions API", "Model", "LLM calls", "LLM cost (USD)", "Latency (ms)",
                     "Evidence ids", "Uncertainties", "Sources", "Cloud families"):
        assert expected in names


def _outcome(company: str = "Y Corp") -> LeadOutcome:
    lead = Lead("Ann", "ann@x.com", company, "https://y.com")
    fit = FitResult(70, "HIGH", 25, 35, 10, "r", cloud_spend_band="HIGH", cloud_spend_reasoning="x")
    cloud_usage = CloudUsageAssessment(
        state="LIKELY", families_present={"engineering_footprint": {"count": 2, "max_strength": "STRONG"}})
    research = Research(summary="s", sources=["https://y.com"], cloud_usage=cloud_usage,
                        evidence_ids=["ev-001", "ev-002"], uncertainties=["no headcount source"])
    return LeadOutcome(lead, research,
                       ComplianceResult(False, "clear", [], "ok", sanctions_status="none"), fit, True,
                       llm_calls=3, llm_cost_usd=0.00042, llm_latency_ms=1200,
                       llm_models="deepseek/deepseek-v4-flash-0731:free", llm_hops=2,
                       run_id="11111111-2222-3333-4444-555555555555", evidence_count=4,
                       provenance_status="ok", provenance_head="4:abc123")


def test_leads_sheet_gets_sales_values(tmp_path: Path):
    p = tmp_path / "t.xlsx"
    row = write_row(_outcome(), p)
    wb = load_workbook(p)
    ws = wb["Leads"]
    idx = {c.value: i + 1 for i, c in enumerate(ws[1])}
    assert ws.cell(row, idx["Fit score"]).value == 70
    assert ws.cell(row, idx["Fit confidence"]).value == "HIGH"
    assert ws.cell(row, idx["Spend prior (relative)"]).value == "HIGH"
    assert ws.cell(row, idx["Compliance"]).value == "clear"
    assert ws.cell(row, idx["Cloud usage"]).value == "LIKELY"


def test_technical_sheet_gets_run_and_telemetry_values(tmp_path: Path):
    p = tmp_path / "t.xlsx"
    row = write_row(_outcome(), p)
    wb = load_workbook(p)
    ws = wb["Technical"]
    idx = {c.value: i + 1 for i, c in enumerate(ws[1])}
    assert ws.cell(row, idx["Run ID"]).value == "11111111-2222-3333-4444-555555555555"
    assert ws.cell(row, idx["Evidence count"]).value == 4
    assert ws.cell(row, idx["Provenance status"]).value == "ok"
    assert ws.cell(row, idx["Provenance head"]).value == "4:abc123"
    assert ws.cell(row, idx["Sanctions API"]).value == "none"
    assert ws.cell(row, idx["Model"]).value == "deepseek/deepseek-v4-flash-0731:free"
    assert ws.cell(row, idx["LLM calls"]).value == 3
    assert ws.cell(row, idx["LLM cost (USD)"]).value == 0.00042
    assert ws.cell(row, idx["Latency (ms)"]).value == 1200
    assert "engineering_footprint" in ws.cell(row, idx["Cloud families"]).value
    assert ws.cell(row, idx["Evidence ids"]).value == "ev-001, ev-002"
    assert ws.cell(row, idx["Uncertainties"]).value == "no headcount source"


def test_rerun_updates_both_sheets_in_place(tmp_path: Path):
    p = tmp_path / "t.xlsx"
    assert write_row(_outcome(), p) == 2
    assert write_row(_outcome(), p) == 2  # same company updates, doesn't append
    wb = load_workbook(p)
    assert wb["Leads"].max_row == 2
    assert wb["Technical"].max_row == 2


def test_contact_actionable_column_is_conservative(tmp_path: Path):
    """Fix #8: only an explicit "valid" contact_quality reads as actionable -
    "unverified" (the default) must not collapse into "yes"."""
    from leadscout.models import Lead as _Lead

    p = tmp_path / "t.xlsx"
    unverified = _outcome("Unverified Co")
    write_row(unverified, p)
    valid = _outcome("Valid Co")
    valid.research.contact_quality = "valid"
    valid.lead = _Lead("Ann", "ann@x.com", "Valid Co", "https://y.com")
    write_row(valid, p)

    wb = load_workbook(p)
    ws = wb["Leads"]
    idx = {c.value: i + 1 for i, c in enumerate(ws[1])}
    rows = {ws.cell(i, 2).value: i for i in range(2, ws.max_row + 1)}
    assert ws.cell(rows["Unverified Co"], idx["Contact actionable"]).value == "no"
    assert ws.cell(rows["Valid Co"], idx["Contact actionable"]).value == "yes"


def test_infrastructure_column_summarises_footprints(tmp_path: Path):
    from leadscout.models import InfrastructureFootprint

    p = tmp_path / "t.xlsx"
    outcome = _outcome()
    outcome.research.infrastructure = [
        InfrastructureFootprint(scope="website", category="MANAGED_HOSTING", provider=None,
                                network_holder="WEBSUPPORT-AS", confidence="MEDIUM", evidence_ids=["ev-003"]),
    ]
    row = write_row(outcome, p)
    wb = load_workbook(p)
    ws = wb["Leads"]
    idx = {c.value: i + 1 for i, c in enumerate(ws[1])}
    value = ws.cell(row, idx["Infrastructure"]).value
    assert "MANAGED_HOSTING" in value
    assert "WEBSUPPORT-AS" in value


def test_rewriting_a_lead_never_keeps_a_stale_value(tmp_path: Path):
    """openpyxl's ws.cell(value=None) keeps the old value; an update must overwrite it."""
    p = tmp_path / "t.xlsx"
    first = _outcome()
    first.research.estimated_employees = 1000
    write_row(first, p)
    second = _outcome()
    second.lead.company_size_band = "201-1000"
    row = write_row(second, p)
    ws = load_workbook(p)["Leads"]
    idx = {c.value: i + 1 for i, c in enumerate(ws[1])}
    assert ws.cell(row, idx["Employees (est.)"]).value == "~500 (self-reported band 201-1000)"
    third = _outcome()
    row = write_row(third, p)
    assert load_workbook(p)["Leads"].cell(row, idx["Employees (est.)"]).value is None


def test_next_action_routes_work_and_compliance_outranks_fit():
    from leadscout.notify import next_action
    ok = _outcome()
    ok.research.contact_quality = "valid"
    assert next_action(ok).startswith("CONTACT NOW")
    bad_contact = _outcome()
    bad_contact.research.contact_quality = "invalid"
    assert next_action(bad_contact).startswith("WORK THE ACCOUNT")
    blocked = _outcome()
    blocked.compliance.status = "blocked"  # fit 70 / sales_ready True must not matter
    assert next_action(blocked).startswith("DO NOT ENGAGE")
    review = _outcome()
    review.compliance.status = "review"
    assert next_action(review).startswith("HOLD")


def test_an_older_column_layout_is_migrated_by_name_not_by_position(tmp_path: Path):
    from openpyxl import Workbook
    p = tmp_path / "old.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(["Received (UTC)", "Company", "Website", "Fit score", "Compliance"])
    ws.append(["2026-01-01 00:00", "Old Co", "https://old.example", 42, "clear"])
    wb.save(p)
    write_row(_outcome("New Co"), p)
    ws = load_workbook(p)["Leads"]
    idx = {c.value: i + 1 for i, c in enumerate(ws[1])}
    assert [c.value for c in ws[1]][:3] == ["Received (UTC)", "Company", "Next action"]
    assert ws.cell(2, idx["Company"]).value == "Old Co"
    assert ws.cell(2, idx["Fit score"]).value == 42 and ws.cell(2, idx["Website"]).value == "https://old.example"
    assert ws.cell(3, idx["Company"]).value == "New Co" and ws.cell(3, idx["Next action"]).value


def test_unknown_size_with_cloud_evidence_is_a_question_not_a_rejection():
    from leadscout.notify import next_action
    o = _outcome()
    o.sales_ready = False
    o.research.estimated_employees = None
    o.research.employees_hint = "web search suggests 1920 - web search: quoted snippet (ev-9); entity not verified"
    action = next_action(o)  # _outcome() has cloud LIKELY and no size band
    assert action.startswith("CONFIRM SIZE") and "1920" in action and "unverified" in action
