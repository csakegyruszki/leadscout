"""Excel tracker a sales rep can open. One row per lead; re-running a lead updates its row.

Two sheets, split by audience: "Leads" is what a sales rep opens (nothing about run
ids, evidence counts or ledger internals) and "Technical" is what an engineer or
auditor opens to trace a row back to its provenance run and raw telemetry.
"""
from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import LeadOutcome
from .notify import next_action
from .profile_config import fit_config

# Skim-first order (REVIEW-D): what to do and why in the first columns a rep sees; contact detail
# next; the long narrative columns last, so 50 rows can be scanned without horizontal scrolling.
LEADS_COLUMNS = [
    ("Received (UTC)", 18), ("Company", 24), ("Next action", 46), ("Sales-ready", 11),
    ("Contact actionable", 12), ("Fit score", 10), ("Fit confidence", 12), ("Compliance", 12),
    ("Cloud usage", 14), ("Employees (est.)", 16), ("Industry", 18), ("HQ country", 14),
    ("Contact", 20), ("Email", 26), ("Website", 28), ("Contact quality", 14), ("Spend prior (relative)", 16),
    ("Summary", 60), ("Compliance reasoning", 50), ("Fit reasoning", 50),
    # Fix #8: ACCOUNT (fit+compliance, already the "Sales-ready" column above) and
    # CONTACT (contact quality) are two different questions - kept as separate
    # columns instead of one merged flag. "Infrastructure" mirrors notify.py's
    # per-footprint block, one line per footprint joined with "; ".
    ("Infrastructure", 60),
]
TECHNICAL_COLUMNS = [
    ("Run ID", 38), ("Company", 24), ("Evidence count", 14),
    ("Provenance status", 16), ("Provenance head", 24), ("Sanctions API", 16),
    ("Model", 30), ("LLM calls", 10), ("LLM cost (USD)", 14), ("Latency (ms)", 12),
    ("Evidence ids", 30), ("Uncertainties", 40), ("Sources", 40), ("Cloud families", 40),
    ("HQ source", 12),
]
_LEADS_COMPLIANCE_COL = [n for n, _ in LEADS_COLUMNS].index("Compliance") + 1  # 1-indexed
_TECH_LLM_COST_COL = 9  # 1-indexed; keep in sync with TECHNICAL_COLUMNS above
_FILL = {"blocked": "F8CBAD", "review": "FFE699", "clear": "C6E0B4"}


def _new_sheet(wb: Workbook, title: str, columns: list[tuple[str, int]]):
    ws = wb.create_sheet(title) if title not in wb.sheetnames else wb[title]
    if ws.max_row == 1 and ws.cell(1, 1).value is None:
        for i, (name, width) in enumerate(columns, 1):
            c = ws.cell(row=1, column=i, value=name)
            c.font = Font(bold=True)
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    return ws


def _migrate_header(ws, columns: list[tuple[str, int]]) -> None:
    """An existing sheet written with an older column order is rewritten BY COLUMN NAME, so a
    changed layout never silently misaligns a rep's existing rows."""
    expected = [n for n, _ in columns]
    current = [c.value for c in ws[1]]
    if current[:len(expected)] == expected or ws.cell(1, 1).value is None:
        return
    rows = [dict(zip(current, [c.value for c in r], strict=False)) for r in ws.iter_rows(min_row=2)]
    fills = [ws.cell(i, current.index("Compliance") + 1).fill.fgColor.rgb if "Compliance" in current else None
             for i in range(2, ws.max_row + 1)]
    ws.delete_rows(1, ws.max_row)
    for i, (name, width) in enumerate(columns, 1):
        ws.cell(row=1, column=i, value=name).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    for ri, (row, fill) in enumerate(zip(rows, fills, strict=True), 2):
        for ci, name in enumerate(expected, 1):
            cell = ws.cell(row=ri, column=ci)
            cell.value = row.get(name)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if fill and "Compliance" in expected:
            ws.cell(ri, expected.index("Compliance") + 1).fill = PatternFill("solid", fgColor=fill[-6:])


def _ensure(path: Path) -> Workbook:
    if path.exists():
        wb = load_workbook(path)
        if "Leads" in wb.sheetnames:
            _migrate_header(wb["Leads"], LEADS_COLUMNS)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    _new_sheet(wb, "Leads", LEADS_COLUMNS)
    _new_sheet(wb, "Technical", TECHNICAL_COLUMNS)
    return wb


def _find_row(ws, key: str) -> int:
    row = next((i for i in range(2, ws.max_row + 1) if str(ws.cell(i, 2).value or "").strip().lower() == key), None)
    return row if row is not None else ws.max_row + 1


def _employees_cell(r, lead) -> object:
    """Same fallback as the notification header, so the two never disagree."""
    if r.estimated_employees:
        return r.estimated_employees
    bands = fit_config().band_employees()
    if lead.company_size_band in bands:
        return f"~{bands[lead.company_size_band]} (self-reported band {lead.company_size_band})"
    return None


def write_row(outcome: LeadOutcome, path: Path) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = _ensure(path)
    lead, r, c, f = outcome.lead, outcome.research, outcome.compliance, outcome.fit
    key = lead.company.strip().lower()

    # Fix #8: same conservative rule as notify.py - only an explicit "valid"
    # contact_quality counts as actionable; "unverified"/"risky"/"invalid" all
    # render "no" (unknown must not collapse into valid).
    contact_actionable = "yes" if r.contact_quality == "valid" else "no"
    infra_lines = []
    for fp in r.infrastructure:
        who = fp.provider or (f"network holder {fp.network_holder} (operator not established)"
                              if fp.network_holder else "unknown")
        infra_lines.append(f"{fp.scope}/{fp.category}/{who}/{fp.confidence}")
    infrastructure_summary = "; ".join(infra_lines)

    leads_ws = wb["Leads"]
    row = _find_row(leads_ws, key)
    by_name = {
        "Received (UTC)": datetime.now(UTC).strftime("%Y-%m-%d %H:%M"), "Company": lead.company,
        "Next action": next_action(outcome), "Sales-ready": "yes" if outcome.sales_ready else "no",
        "Contact actionable": contact_actionable, "Fit score": f.score, "Fit confidence": f.confidence,
        "Compliance": c.status, "Cloud usage": r.cloud_usage.state, "Employees (est.)": _employees_cell(r, lead),
        "Industry": r.industry, "HQ country": r.headquarters_country, "Contact": lead.name,
        "Email": lead.email, "Website": lead.website, "Contact quality": r.contact_quality,
        "Spend prior (relative)": f.cloud_spend_band, "Summary": r.summary,
        "Compliance reasoning": c.reasoning, "Fit reasoning": f.reasoning,
        "Infrastructure": infrastructure_summary,
    }
    leads_values = [by_name[name] for name, _ in LEADS_COLUMNS]  # KeyError = a column without a value
    for i, v in enumerate(leads_values, 1):
        # openpyxl's ws.cell(value=None) silently keeps the OLD value - assign explicitly so an
        # updated lead never shows a stale figure from an earlier run.
        cell = leads_ws.cell(row=row, column=i)
        cell.value = v
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    leads_ws.cell(row=row, column=_LEADS_COMPLIANCE_COL).fill = PatternFill(
        "solid", fgColor=_FILL.get(c.status, "FFFFFF"))

    tech_ws = wb["Technical"]
    tech_row = _find_row(tech_ws, key)
    cloud_families = ", ".join(
        f"{fam} ({info['max_strength']}, x{info['count']})" for fam, info in r.cloud_usage.families_present.items())
    tech_values: list[object] = [
        outcome.run_id, lead.company, outcome.evidence_count,
        outcome.provenance_status, outcome.provenance_head, c.sanctions_status,
        outcome.llm_models, outcome.llm_calls, round(outcome.llm_cost_usd, 6), outcome.llm_latency_ms,
        ", ".join(r.evidence_ids), "; ".join(r.uncertainties), "\n".join(r.sources), cloud_families,
        r.hq_source,
    ]
    # Named `tv`, not `v` (the "Leads" loop above already binds `v` to a narrower
    # inferred type from `leads_values` - mypy unifies a variable's type across the
    # whole function body, not per-loop, so reusing `v` here with `tech_values`'
    # wider `list[object]` was flagged as re-assigning `v` to an incompatible type).
    for i, tv in enumerate(tech_values, 1):
        # openpyxl's own stubs type Cell.value as int|str|float|None; we intentionally
        # write whatever the outcome carries (a run id, a joined string, a count).
        cell = tech_ws.cell(row=tech_row, column=i)
        cell.value = tv  # type: ignore[assignment]
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    tech_ws.cell(row=tech_row, column=_TECH_LLM_COST_COL).number_format = "0.000000"

    wb.save(path)
    return row


# Backward-compatible alias: earlier tests/tooling import COLUMNS for the sales sheet.
COLUMNS = LEADS_COLUMNS
